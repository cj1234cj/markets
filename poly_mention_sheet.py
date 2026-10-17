"""Weekly spreadsheet of every Polymarket earnings-call mention strike resolving in the next 2 weeks.

Polymarket's Gamma API is public (no login). Pulls open events tagged `earnings-calls`
(plus any `mention-markets` event whose title says "earnings call"), keeps those whose
end date falls within --days of today, and writes one Excel file:

  Strikes  - one row per word, prices in cents
  Events   - one row per earnings call (end date, strike count, volume)

Polymarket's end date is a resolve-by deadline (the call happens on or before it),
not the call date itself.

Output: exports/polymarket_mentions_<YYYY-MM-DD>.xlsx (git-ignored)

  python poly_mention_sheet.py            # next 14 days
  python poly_mention_sheet.py --days 21
"""
import argparse
import datetime as dt
import json
import pathlib
import re
import time

from openpyxl import Workbook

from mention_sheet import OUT_DIR, autosize, get_json, style_header

GAMMA = "https://gamma-api.polymarket.com"


def json_list(x):
    """Polymarket returns some lists as JSON-encoded strings."""
    if isinstance(x, list):
        return x
    try:
        return json.loads(x) if x else []
    except (TypeError, ValueError):
        return []


def cents(v):
    try:
        return round(float(v) * 100, 1) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def num(v):
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def end_day(s):
    """Polymarket end dates are UTC; 03:59Z is midnight US Eastern, so read them in ET."""
    if not s:
        return None
    t = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
    return (t - dt.timedelta(hours=4)).date()


def company_name(e):
    t = re.match(r"What will (.+?) say", e.get("title", ""))
    return t.group(1) if t else e.get("title", "")


def fetch_tag(slug):
    events, offset = [], 0
    while True:
        page = get_json(f"{GAMMA}/events", {"tag_slug": slug, "closed": "false",
                                            "limit": 100, "offset": offset})
        events += page
        if len(page) < 100:
            return events
        offset += 100
        time.sleep(0.1)


def fetch_events():
    seen, out = set(), []
    for e in fetch_tag("earnings-calls") + fetch_tag("mention-markets"):
        if e["id"] in seen:
            continue
        tags = {t.get("slug") for t in e.get("tags") or []}
        if "earnings-calls" in tags or "earnings call" in e.get("title", "").lower():
            seen.add(e["id"])
            out.append(e)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--days", type=int, default=14)
    ap.add_argument("--out", help="output path (default exports/polymarket_mentions_<date>.xlsx)")
    args = ap.parse_args()

    today = dt.date.today()
    horizon = today + dt.timedelta(days=args.days)

    strikes, calls = [], []
    for e in fetch_events():
        day = end_day(e.get("endDate"))
        if day is None or not (today <= day <= horizon):
            continue
        company = company_name(e)
        url = f"https://polymarket.com/event/{e.get('slug', '')}"
        markets = [m for m in e.get("markets", []) if m.get("active") and not m.get("closed")]
        for m in markets:
            yes_bid, yes_ask = cents(m.get("bestBid")), cents(m.get("bestAsk"))
            prices = json_list(m.get("outcomePrices"))
            strikes.append([
                day, company, m.get("groupItemTitle") or m.get("question") or "",
                yes_bid, yes_ask,
                round(100 - yes_ask, 1) if yes_ask is not None else None,  # no bid
                round(100 - yes_bid, 1) if yes_bid is not None else None,  # no ask
                cents(prices[0]) if prices else None, cents(m.get("lastTradePrice")),
                round(yes_ask - yes_bid, 1) if yes_bid is not None and yes_ask is not None else None,
                num(m.get("volumeNum") or m.get("volume")), num(m.get("volume24hr")),
                num(m.get("liquidityNum") or m.get("liquidity")),
                m.get("question"), url,
            ])
        calls.append([day, company, e.get("title"), len(markets),
                      sum(num(m.get("volumeNum") or m.get("volume")) or 0 for m in markets), url])

    strikes.sort(key=lambda r: (r[0], r[1], r[2].lower()))
    calls.sort(key=lambda r: (r[0], r[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Strikes"
    ws.append(["Resolves by", "Company", "Word", "Yes bid ¢", "Yes ask ¢", "No bid ¢", "No ask ¢",
               "Mid ¢", "Last ¢", "Spread ¢", "Volume $", "Volume 24h $", "Liquidity $",
               "Question", "Polymarket link"])
    for r in strikes:
        ws.append(r)
    style_header(ws)
    autosize(ws, {"N": 60, "O": 45})

    ws2 = wb.create_sheet("Events")
    ws2.append(["Resolves by", "Company", "Event", "Strikes", "Total volume $", "Polymarket link"])
    for r in calls:
        ws2.append(r)
    style_header(ws2)
    autosize(ws2, {"F": 45})

    for sheet in (ws, ws2):
        for (cell,) in sheet.iter_rows(min_row=2, max_col=1):
            cell.number_format = "ddd mmm d"

    out = pathlib.Path(args.out) if args.out else OUT_DIR / f"polymarket_mentions_{today.isoformat()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{len(strikes)} strikes across {len(calls)} earnings calls "
          f"({today} to {horizon}) -> {out}")


if __name__ == "__main__":
    main()
