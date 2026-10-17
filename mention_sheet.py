"""Weekly spreadsheet of every Kalshi earnings-call mention strike deciding in the next 2 weeks.

Kalshi market data is public (no login). Pulls all open KXEARNINGSMENTION* events,
keeps those whose call date (from the event ticker, e.g. -26OCT29) falls within
--days of today, and writes one Excel file:

  Strikes  - one row per word, prices in cents, contracts buyable within 3¢ of the ask
  Events   - one row per earnings call (date, strike count, volume)

Output: exports/earnings_mentions_<YYYY-MM-DD>.xlsx (git-ignored)

  python mention_sheet.py            # next 14 days
  python mention_sheet.py --days 21
  python mention_sheet.py --out site/data/earnings_mentions.xlsx --json site/data/kalshi_mentions.json   # dashboard (CI)
"""
import argparse
import datetime as dt
import json
import pathlib
import re
import time

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KALSHI = "https://api.elections.kalshi.com/trade-api/v2"
OUT_DIR = pathlib.Path(__file__).parent / "exports"
TICKER_DATE = re.compile(r"-(\d{2}[A-Z]{3}\d{2})$")


def get_json(url, params, tries=4):
    for i in range(tries):
        try:
            r = requests.get(url, params=params, timeout=30)
            if r.status_code == 429:
                time.sleep(2 ** i)
                continue
            r.raise_for_status()
            return r.json()
        except requests.RequestException:
            if i == tries - 1:
                raise
            time.sleep(2 ** i)
    raise RuntimeError(f"gave up on {url}")


def cents(m, key):
    """Kalshi prices: newer `<key>_dollars` string, else legacy integer cents."""
    d = m.get(f"{key}_dollars")
    if d not in (None, ""):
        return round(float(d) * 100)
    v = m.get(key)
    return v if isinstance(v, (int, float)) else None


def qty(m, key):
    v = m.get(f"{key}_fp", m.get(key))
    return float(v) if v not in (None, "") else None


def event_date(ticker):
    m = TICKER_DATE.search(ticker)
    if not m:
        return None
    try:
        return dt.datetime.strptime(m.group(1), "%y%b%d").date()
    except ValueError:
        return None


def company_name(e):
    t = re.match(r"What will (.+?) say", e.get("title", ""))
    if t:
        return t.group(1)
    src = e.get("settlement_sources") or []
    if src and src[0].get("name"):
        return re.sub(r"\s+Investor (Events|Relations)$", "", src[0]["name"])
    return e.get("title", "")


def fetch_events():
    events, cursor = [], None
    while True:
        params = {"status": "open", "with_nested_markets": "true", "limit": 200}
        if cursor:
            params["cursor"] = cursor
        data = get_json(f"{KALSHI}/events", params)
        events += [e for e in data.get("events", [])
                   if e.get("event_ticker", "").startswith("KXEARNINGSMENTION")]
        cursor = data.get("cursor")
        if not cursor:
            return events
        time.sleep(0.1)


FEE_RATE = 0.07  # Kalshi taker fee: 0.07 x C x P x (1-P), rounded up to the cent per order


def fee(price_c):
    """Taker fee per contract in cents at a price in cents (before the per-order round-up)."""
    p = price_c / 100
    return FEE_RATE * p * (1 - p) * 100


def all_in(price_c):
    return None if price_c is None else round(price_c + fee(price_c), 2)


def payout(cost_c):
    """Profit per $1 staked if the bet wins, in %."""
    return None if not cost_c else round((100 / cost_c - 1) * 100, 1)


def asks_from(opposite_bids):
    """Kalshi books list bids only. Buying YES fills against NO bids (a NO bid at p is a YES ask
    at 100 - p), so turn the opposite side's bids into asks, cheapest first: [(price ¢, qty)]."""
    return sorted((round(100 - float(p) * 100, 2), float(q)) for p, q in opposite_bids or [])


def depth_within(asks, window=3):
    """Contracts you can buy within `window`¢ of the best ask."""
    if not asks:
        return 0
    return round(sum(q for p, q in asks if p <= asks[0][0] + window + 1e-9))


def fill(asks, budget):
    """Walk the book spending `budget` dollars, fees included.
    Returns (avg price ¢ before fees, avg all-in ¢, contracts, filled in full)."""
    left, n, spent = budget * 100, 0.0, 0.0
    for p, q in asks:
        cost = p + fee(p)
        take = min(q, left / cost)
        n += take
        spent += take * p
        left -= take * cost
        if left <= 1e-9:
            break
    if not n:
        return None, None, 0, False
    return round(spent / n, 2), round((budget * 100 - left) / n, 2), round(n), left <= 1e-6


def fetch_books(tickers):
    """{ticker: (yes asks, no asks)} from the batch orderbook endpoint, 100 tickers per call."""
    out = {}
    for i in range(0, len(tickers), 100):
        params = [("tickers", t) for t in tickers[i:i + 100]]
        for ob in get_json(f"{KALSHI}/markets/orderbooks", params).get("orderbooks", []):
            book = ob.get("orderbook_fp") or {}
            out[ob["ticker"]] = (asks_from(book.get("no_dollars")), asks_from(book.get("yes_dollars")))
        time.sleep(0.1)
    return out


def autosize(ws, widths=None):
    for i, col in enumerate(ws.columns, 1):
        letter = get_column_letter(i)
        if widths and letter in widths:
            ws.column_dimensions[letter].width = widths[letter]
            continue
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[letter].width = min(max(10, longest + 2), 50)


def write_json(path, source, today, horizon, strikes, cols=None):
    """Strikes grouped by call, for the readable mentions page. Rows are lists that start with
    date, company and end with the event link, picked into JSON by `cols` ({key: index}), or
    dicts with "date", "company", "url" keys (cols=None)."""
    events = {}
    for r in strikes:
        if cols is None:
            key, s = (r["date"], r["company"], r["url"]), {k: v for k, v in r.items()
                                                         if k not in ("date", "company", "url")}
        else:
            key, s = (r[0], r[1], r[-1]), {k: r[i] for k, i in cols.items()}
        ev = events.setdefault(key, {"date": key[0].isoformat(), "company": key[1], "url": key[2],
                                     "strikes": []})
        ev["strikes"].append(s)
    for ev in events.values():
        ev["volume"] = sum(s.get("volume") or 0 for s in ev["strikes"])
    payload = {"source": source, "from": today.isoformat(), "to": horizon.isoformat(),
               "generated": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
               "events": list(events.values())}
    path = pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")


def style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


RED, YELLOW, GREEN = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFE699", "C6EFCE"))


def depth_fill(n):
    """Thin books stand out: red under 100 contracts, yellow 100-300, green above 300."""
    if n is None:
        return None
    return RED if n < 100 else YELLOW if n <= 300 else GREEN


SIZES = (50, 100, 250, 500)  # order sizes on the Fills sheet


def sized_fill(asks, budget):
    """Fill for one order size. When the whole visible book can't take it, there's no average
    to quote: report the contracts available instead (a limit order is the only route)."""
    avg, avg_all_in, n, full = fill(asks, budget)
    if not full:
        return {"full": False, "n": n, "fill": None, "fill_all_in": None, "payout": None,
                "text": f"partial fill: {n:,} contracts available" if n else "no offers in book"}
    return {"full": True, "n": n, "fill": avg, "fill_all_in": avg_all_in,
            "payout": payout(avg_all_in), "text": None}


def side(asks, ask_c, budget):
    """Pre-trade numbers for buying one side: all-in cost, payout, depth, fills by order size."""
    cost = all_in(ask_c)
    return {"ask": ask_c, "all_in": cost, "payout": payout(cost), "depth": depth_within(asks),
            "fills": {b: sized_fill(asks, b) for b in sorted({*SIZES, budget})},
            "book": [[p, round(q, 2)] for p, q in asks]}


def fill_cell(f):
    return f["text"] or f["fill_all_in"]


def side_columns(name, k, budget):
    usd = f"${budget:g}"
    return [
        (f"{name} ask ¢", lambda r: r[k]["ask"]),
        (f"{name} all-in ¢ (ask + fee)", lambda r: r[k]["all_in"]),
        (f"{name} payout % at ask (profit per $1)", lambda r: r[k]["payout"]),
        (f"{name} contracts ≤3¢ of ask", lambda r: r[k]["depth"]),
        (f"{name} {usd} fill ¢ (with fees)", lambda r: fill_cell(r[k]["fills"][budget])),
        (f"{name} {usd} fill ¢ (before fees)", lambda r: r[k]["fills"][budget]["fill"]),
        (f"{name} payout % at {usd} fill", lambda r: r[k]["fills"][budget]["payout"]),
    ]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--days", type=int, default=14)
    ap.add_argument("--order", type=float, default=100, help="order size in $ for the fill estimate")
    ap.add_argument("--out", help="output path (default exports/earnings_mentions_<date>.xlsx)")
    ap.add_argument("--json", help="also write strikes grouped by call, for the web page")
    args = ap.parse_args()

    today = dt.date.today()
    horizon = today + dt.timedelta(days=args.days)
    events = [e for e in fetch_events()
              if (d := event_date(e["event_ticker"])) and today <= d <= horizon]
    books = fetch_books([m["ticker"] for e in events for m in e.get("markets", [])])

    strikes, calls = [], []
    for e in events:
        day = event_date(e["event_ticker"])
        company = company_name(e)
        url = f"https://kalshi.com/markets/{e.get('series_ticker', '').lower()}"
        markets = [m for m in e.get("markets", []) if m.get("status") in ("active", "open", None)]
        for m in markets:
            yes_bid, yes_ask = cents(m, "yes_bid"), cents(m, "yes_ask")
            yes_asks, no_asks = books.get(m.get("ticker"), ([], []))
            strikes.append({
                "date": day, "company": company, "url": url,
                "word": m.get("yes_sub_title") or m.get("subtitle") or "",
                "yes": side(yes_asks, yes_ask, args.order),
                "no": side(no_asks, cents(m, "no_ask"), args.order),
                "yes_bid": yes_bid, "no_bid": cents(m, "no_bid"), "last": cents(m, "last_price"),
                "spread": (yes_ask - yes_bid) if yes_bid is not None and yes_ask is not None else None,
                "volume": qty(m, "volume"), "volume_24h": qty(m, "volume_24h"),
                "open_interest": qty(m, "open_interest"), "ticker": m.get("ticker"),
            })
        calls.append([day, company, e["event_ticker"], len(markets),
                      sum(qty(m, "volume") or 0 for m in markets), url])

    strikes.sort(key=lambda r: (r["date"], r["company"], r["word"].lower()))
    calls.sort(key=lambda r: (r[0], r[1]))

    usd = f"${args.order:g}"
    columns = ([("Call date", lambda r: r["date"]), ("Company", lambda r: r["company"]),
                ("Word", lambda r: r["word"])]
               + side_columns("Yes", "yes", args.order) + side_columns("No", "no", args.order)
               + [("Yes bid ¢", lambda r: r["yes_bid"]), ("No bid ¢", lambda r: r["no_bid"]),
                  ("Last ¢", lambda r: r["last"]), ("Spread ¢", lambda r: r["spread"]),
                  ("Volume", lambda r: r["volume"]), ("Volume 24h", lambda r: r["volume_24h"]),
                  ("Open interest", lambda r: r["open_interest"]), ("Ticker", lambda r: r["ticker"]),
                  ("Kalshi link", lambda r: r["url"])])

    wb = Workbook()
    ws = wb.active
    ws.title = "Strikes"
    ws.append([h for h, _ in columns])
    for r in strikes:
        ws.append([get(r) for _, get in columns])
    # colour flags: payouts under 33% red; depth red/yellow/green; a partial fill red
    col = {h: i + 1 for i, (h, _) in enumerate(columns)}
    for row, r in enumerate(strikes, 2):
        for name, k in (("Yes", "yes"), ("No", "no")):
            sd, f = r[k], r[k]["fills"][args.order]
            if sd["payout"] is not None and sd["payout"] < 33:
                ws.cell(row, col[f"{name} payout % at ask (profit per $1)"]).fill = RED
            if dfill := depth_fill(sd["depth"]):
                ws.cell(row, col[f"{name} contracts ≤3¢ of ask"]).fill = dfill
            if not f["full"]:
                ws.cell(row, col[f"{name} {usd} fill ¢ (with fees)"]).fill = RED
            elif f["payout"] is not None and f["payout"] < 33:
                ws.cell(row, col[f"{name} payout % at {usd} fill"]).fill = RED
    style_header(ws)
    for c in ws[1]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 45
    autosize(ws, {get_column_letter(i): 11 for i in range(4, 18)}
             | {get_column_letter(col[f"{n} {usd} fill ¢ (with fees)"]): 16 for n in ("Yes", "No")}
             | {get_column_letter(len(columns)): 45})

    # every order size side by side: fill with fees and the payout you'd actually get
    wf = wb.create_sheet("Fills")
    wf.append(["Call date", "Company", "Word"]
              + [f"{n} ${b} {what}" for n in ("Yes", "No") for b in SIZES
                 for what in ("fill ¢ (with fees)", "payout %")])
    for row, r in enumerate(strikes, 2):
        cells = [r["date"], r["company"], r["word"]]
        for k in ("yes", "no"):
            for b in SIZES:
                f = r[k]["fills"][b]
                cells += [fill_cell(f), f["payout"]]
        wf.append(cells)
        for i, k in enumerate(("yes", "no")):
            for j, b in enumerate(SIZES):
                f, c0 = r[k]["fills"][b], 4 + (i * len(SIZES) + j) * 2
                if not f["full"]:
                    wf.cell(row, c0).fill = RED
                elif f["payout"] is not None and f["payout"] < 33:
                    wf.cell(row, c0 + 1).fill = RED
    style_header(wf)
    for c in wf[1]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wf.row_dimensions[1].height = 45
    autosize(wf, {get_column_letter(i): 13 for i in range(4, 4 + 4 * len(SIZES))})

    try:  # CI runs in UTC; show Eastern like the dashboard
        from zoneinfo import ZoneInfo
        now = dt.datetime.now(ZoneInfo("America/New_York"))
    except Exception:
        now = dt.datetime.now().astimezone()
    updated = now.strftime("%a %b %d %Y, %I:%M %p %Z")
    ws3 = wb.create_sheet("Key")
    for line in [
        ["Updated", f"{updated} (prices and order books as of this time; they move fast)"],
        ["All-in ¢", "Ask plus Kalshi's taker fee, 0.07 x P x (1 - P). Kalshi rounds the fee up to the "
                     "next cent per order, so a tiny order costs up to 1¢ more."],
        ["Payout % at ask", "Profit per $1 if the bet wins, at the all-in ask: 100 / all-in - 1. "
                            "Red under 33% (cost above about 75¢). Overstates the payout when slippage is "
                            "large, so check payout at fill."],
        ["Contracts ≤3¢", "Contracts you can buy within 3¢ of the ask. Red under 100, yellow 100-300, "
                          "green above 300."],
        [f"{usd} fill ¢", f"Walks the full order book spending {usd}, fees included in the {usd}. "
                          f"'With fees' is comparable to all-in; 'before fees' is the raw book price."],
        ["Payout at fill", "Profit per $1 at the with-fees fill price: what you'd actually get. Red under 33%."],
        ["Partial fill", "The whole visible book can't take the order: shows the contracts available "
                         "instead of a price (red). A limit order is the only route."],
        ["Fills tab", f"The same fill and payout for ${', $'.join(map(str, SIZES))} orders. "
                      f"The web page takes any size."],
    ]:
        ws3.append(line)
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 110

    ws2 = wb.create_sheet("Events")
    ws2.append(["Call date", "Company", "Event ticker", "Strikes", "Total volume", "Kalshi link"])
    for r in calls:
        ws2.append(r)
    style_header(ws2)
    autosize(ws2, {"F": 45})
    wb.move_sheet("Key", offset=1)

    for sheet in (ws, wf, ws2):
        for (cell,) in sheet.iter_rows(min_row=2, max_col=1):
            cell.number_format = "ddd mmm d"
        sheet.cell(row=1, column=1).number_format = "General"

    out = pathlib.Path(args.out) if args.out else OUT_DIR / f"earnings_mentions_{today.isoformat()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    if args.json:
        write_json(args.json, "kalshi", today, horizon, strikes)
    print(f"{len(strikes)} strikes across {len(calls)} earnings calls "
          f"({today} to {horizon}) -> {out}")


if __name__ == "__main__":
    main()
